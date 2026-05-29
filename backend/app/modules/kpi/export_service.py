"""
Export service for KPI tables — CSV and Excel (openpyxl) formats.
Supported tables: top_customers, top_suppliers, low_stock_alerts.
"""
from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession


def _fmt_minor(v: int) -> str:
    return f"{v / 100:.2f}"


# ---------------------------------------------------------------------------
# Data fetchers
# ---------------------------------------------------------------------------


async def _fetch_top_customers(db: AsyncSession, garage_id: int) -> list[dict[str, Any]]:
    result = await db.execute(
        text(
            """
            SELECT c.full_name AS name,
                   COALESCE(SUM(i.total_minor), 0)::bigint AS revenue_minor,
                   COUNT(i.id)::int AS invoice_count
            FROM customer c
            JOIN invoice i ON i.customer_id = c.id
            WHERE i.garage_id = :g
              AND i.status IN ('issued', 'partial', 'paid')
              AND i.deleted_at IS NULL
              AND c.deleted_at IS NULL
            GROUP BY c.id, c.full_name
            ORDER BY revenue_minor DESC
            LIMIT 10
            """
        ),
        {"g": garage_id},
    )
    return [
        {
            "rank": idx + 1,
            "name": r.name,
            "revenue_eur": _fmt_minor(r.revenue_minor),
            "invoice_count": r.invoice_count,
        }
        for idx, r in enumerate(result.fetchall())
    ]


async def _fetch_top_suppliers(db: AsyncSession, garage_id: int) -> list[dict[str, Any]]:
    result = await db.execute(
        text(
            """
            SELECT s.name,
                   COALESCE(SUM(po.total_ht), 0)::bigint AS spend_minor,
                   COUNT(po.id)::int AS order_count
            FROM supplier s
            JOIN purchase_order po ON po.supplier_id = s.id
            WHERE po.garage_id = :g
              AND po.status NOT IN ('cancelled')
              AND po.deleted_at IS NULL
              AND s.deleted_at IS NULL
            GROUP BY s.id, s.name
            ORDER BY spend_minor DESC
            LIMIT 10
            """
        ),
        {"g": garage_id},
    )
    return [
        {
            "rank": idx + 1,
            "name": r.name,
            "spend_eur": _fmt_minor(r.spend_minor),
            "order_count": r.order_count,
        }
        for idx, r in enumerate(result.fetchall())
    ]


async def _fetch_low_stock_alerts(db: AsyncSession, garage_id: int) -> list[dict[str, Any]]:
    result = await db.execute(
        text(
            """
            SELECT p.sku_code, p.name AS part_name,
                   a.on_hand, a.threshold,
                   COALESCE(a.suggested_reorder_quantity, 0) AS suggested_qty,
                   a.status
            FROM low_stock_alert a
            JOIN part p ON p.id = a.part_id
            WHERE a.garage_id = :g
              AND a.status = 'open'
              AND p.deleted_at IS NULL
            ORDER BY (a.threshold - a.on_hand) DESC
            """
        ),
        {"g": garage_id},
    )
    return [
        {
            "sku": r.sku_code,
            "name": r.part_name,
            "on_hand": r.on_hand,
            "threshold": r.threshold,
            "deficit": r.threshold - r.on_hand,
            "suggested_qty": r.suggested_qty,
            "status": r.status,
        }
        for r in result.fetchall()
    ]


# ---------------------------------------------------------------------------
# CSV export
# ---------------------------------------------------------------------------

_COLUMNS: dict[str, list[tuple[str, str]]] = {
    "top_customers": [
        ("rank", "Rang"),
        ("name", "Client"),
        ("revenue_eur", "CA (€)"),
        ("invoice_count", "Nb factures"),
    ],
    "top_suppliers": [
        ("rank", "Rang"),
        ("name", "Fournisseur"),
        ("spend_eur", "Achats (€)"),
        ("order_count", "Nb commandes"),
    ],
    "low_stock_alerts": [
        ("sku", "Référence"),
        ("name", "Désignation"),
        ("on_hand", "Stock actuel"),
        ("threshold", "Seuil"),
        ("deficit", "Déficit"),
        ("suggested_qty", "Qté suggérée"),
    ],
}


async def export_csv(db: AsyncSession, garage_id: int, table: str) -> str:
    rows = await _fetch_rows(db, garage_id, table)
    columns = _COLUMNS[table]
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([label for _, label in columns])
    for row in rows:
        writer.writerow([row.get(key, "") for key, _ in columns])
    return buf.getvalue()


async def export_xlsx(db: AsyncSession, garage_id: int, table: str) -> bytes:
    rows = await _fetch_rows(db, garage_id, table)
    columns = _COLUMNS[table]

    wb = Workbook()
    ws = wb.active
    ws.title = table

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="1976D2")

    for col_idx, (_, label) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (key, _) in enumerate(columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _fetch_rows(
    db: AsyncSession, garage_id: int, table: str
) -> list[dict[str, Any]]:
    if table == "top_customers":
        return await _fetch_top_customers(db, garage_id)
    if table == "top_suppliers":
        return await _fetch_top_suppliers(db, garage_id)
    if table == "low_stock_alerts":
        return await _fetch_low_stock_alerts(db, garage_id)
    raise ValueError(f"Unknown export table: {table}")
