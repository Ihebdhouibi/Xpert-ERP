"""Customer service — business logic for the customer module.

The service owns:

- Duplicate detection per the M2 merge rule:
  `fiscal_id` AND (`phone_e164` OR `email`) ⇒ merge prompt.
  A `fiscal_id`-only collision is also surfaced as a merge prompt rather than
  letting the DB partial-unique constraint return a 500.
- Optimistic locking on update via the `version` column.
- Soft delete and admin-only restore.
- Audit hooks on every create / update / delete / restore / credit-limit
  change / balance adjustment.
- Atomic balance-event posting (ledger row + customer.outstanding_minor in
  the same transaction).

Service functions flush via the audit writer but do NOT commit — the calling
endpoint owns the transaction (matches the auth/rbac modules).
"""
from __future__ import annotations

import enum
import re
import uuid as uuid_lib
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any, Optional, Union

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.logging import get_logger
from app.modules.audit.service import audit
from app.modules.customer import repository as repo
from app.modules.customer.models import (
    ContactChannel,
    Customer,
    CustomerBalanceEvent,
    CustomerContact,
    CustomerStatus,
    CustomerType,
)
from app.modules.customer.schemas import (
    CreditLimitUpdateRequest,
    CustomerBalanceEventCreateRequest,
    CustomerContactCreate,
    CustomerCreateRequest,
    CustomerQuickCreateRequest,
    CustomerUpdateRequest,
    MergeMatchReason,
)

logger = get_logger("customer")

_SIRET_PATTERN = re.compile(r"^\d{14}$")

# Segment stamped on walk-in customers created from just a name, so they can be
# told apart from fully registered ones (shown as a chip on the customer page).
WALK_IN_SEGMENT = "passager"


# ---------------------------------------------------------------------------
# Exceptions surfaced to the router layer
# ---------------------------------------------------------------------------


class CustomerNotFoundError(Exception):
    """Customer does not exist in the actor's garage."""


class OptimisticLockError(Exception):
    """Submitted `version` does not match the current row."""

    def __init__(self, expected: int, actual: int) -> None:
        super().__init__(
            f"Stale version: client sent {expected}, current is {actual}"
        )
        self.expected = expected
        self.actual = actual


class FiscalIdConflictError(Exception):
    """A fiscal_id update would collide with another live customer."""


class InvalidFiscalIdError(Exception):
    """fiscal_id does not match the required format for the customer's type."""


# ---------------------------------------------------------------------------
# Result types
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class DuplicateMatch:
    """Returned by `create_customer` when the merge rule fires.

    Reason values come from `schemas.MergeMatchReason`.
    """

    existing: Customer
    reason: str


CreateResult = Union[Customer, DuplicateMatch]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _utcnow() -> datetime:
    return datetime.now(tz=timezone.utc)


def _to_audit_dict(obj: Customer) -> dict[str, Any]:
    """Serialize an ORM Customer row to a JSON-safe dict for audit storage."""
    out: dict[str, Any] = {}
    for col in obj.__table__.columns:
        val = getattr(obj, col.name)
        if val is None:
            out[col.name] = None
        elif isinstance(val, datetime):
            out[col.name] = val.isoformat()
        elif isinstance(val, uuid_lib.UUID):
            out[col.name] = str(val)
        elif isinstance(val, enum.Enum):
            out[col.name] = val.value
        else:
            out[col.name] = val
    return out


def _apply_contacts(
    customer: Customer, contacts: list[CustomerContactCreate]
) -> None:
    """Replace the customer's contacts list with the supplied set."""
    customer.contacts.clear()
    for c in contacts:
        customer.contacts.append(
            CustomerContact(
                channel=c.channel,
                value=c.value,
                label=c.label,
                is_primary=c.is_primary,
            )
        )


# ---------------------------------------------------------------------------
# Duplicate detection
# ---------------------------------------------------------------------------


async def _check_duplicate(
    db: AsyncSession,
    garage_id: int,
    payload: CustomerCreateRequest,
) -> Optional[DuplicateMatch]:
    """Apply the M2 merge rule. Returns None if creation can proceed."""
    if not payload.fiscal_id:
        return None  # fiscal_id is the anchor for the merge rule

    existing = await repo.find_by_fiscal_id(db, garage_id, payload.fiscal_id)
    if existing is None:
        return None

    if payload.phone_e164 and existing.phone_e164 == payload.phone_e164:
        return DuplicateMatch(existing, MergeMatchReason.PHONE_AND_FISCAL_ID)
    if payload.email and existing.email and existing.email.lower() == payload.email.lower():
        return DuplicateMatch(existing, MergeMatchReason.EMAIL_AND_FISCAL_ID)

    # fiscal_id collision without contact match: DB partial-unique would
    # reject the INSERT with a generic constraint error; surface it cleanly.
    return DuplicateMatch(existing, MergeMatchReason.FISCAL_ID_ONLY)


# ---------------------------------------------------------------------------
# Create
# ---------------------------------------------------------------------------


async def create_customer(
    db: AsyncSession,
    *,
    garage_id: int,
    actor_user_id: Optional[int],
    payload: CustomerCreateRequest,
    request_id: Optional[str] = None,
    ip_address: Optional[str] = None,
) -> CreateResult:
    """Create a customer or signal a duplicate match.

    Returns a `Customer` on success (caller must commit) or a `DuplicateMatch`
    that the router translates into a 409 merge-prompt response.
    """
    logger.info("create_customer_start", garage_id=garage_id)
    duplicate = await _check_duplicate(db, garage_id, payload)
    if duplicate is not None:
        return duplicate

    new_customer = Customer(
        garage_id=garage_id,
        type=payload.type,
        full_name=payload.full_name,
        company_name=payload.company_name,
        fiscal_id=payload.fiscal_id,
        email=payload.email,
        phone_e164=payload.phone_e164,
        address=payload.address.model_dump(exclude_none=True) if payload.address else None,
        segment=payload.segment,
        credit_limit_minor=payload.credit_limit_minor,
        credit_currency=payload.credit_currency,
        notes=payload.notes,
        created_by_user_id=actor_user_id,
        updated_by_user_id=actor_user_id,
    )
    _apply_contacts(new_customer, payload.contacts)

    db.add(new_customer)
    await db.flush()  # populate id / public_id

    await audit(
        db=db,
        garage_id=garage_id,
        user_id=actor_user_id,
        component="customer",
        action="create",
        entity_type="customer",
        entity_id=new_customer.id,
        new_value=_to_audit_dict(new_customer),
        request_id=request_id,
        ip_address=ip_address,
    )
    logger.info("create_customer_success", customer_id=new_customer.id, garage_id=garage_id)
    return new_customer


async def create_quick_customer(
    db: AsyncSession,
    *,
    garage_id: int,
    actor_user_id: Optional[int],
    payload: CustomerQuickCreateRequest,
    request_id: Optional[str] = None,
    ip_address: Optional[str] = None,
) -> Customer:
    """Create a walk-in customer from just a name (no CIN/SIRET, phone or email).

    Skips the duplicate check on purpose: it is anchored on `fiscal_id`, which a
    walk-in does not have, and two different walk-ins can share a name. Callers
    are expected to let the user pick an existing customer first.
    """
    logger.info("create_quick_customer_start", garage_id=garage_id)
    new_customer = Customer(
        garage_id=garage_id,
        type=CustomerType.individual,
        full_name=payload.full_name,
        segment=WALK_IN_SEGMENT,
        created_by_user_id=actor_user_id,
        updated_by_user_id=actor_user_id,
    )
    db.add(new_customer)
    await db.flush()  # populate id / public_id

    await audit(
        db=db,
        garage_id=garage_id,
        user_id=actor_user_id,
        component="customer",
        action="create",
        entity_type="customer",
        entity_id=new_customer.id,
        new_value=_to_audit_dict(new_customer),
        request_id=request_id,
        ip_address=ip_address,
    )
    logger.info(
        "create_quick_customer_success", customer_id=new_customer.id, garage_id=garage_id
    )
    return new_customer


# ---------------------------------------------------------------------------
# Update
# ---------------------------------------------------------------------------

# Fields the general PATCH endpoint may change. credit_limit_minor and
# credit_currency are excluded — they use a dedicated endpoint with audit
# reason and their own permission.
_UPDATABLE_FIELDS = (
    "full_name",
    "company_name",
    "fiscal_id",
    "email",
    "phone_e164",
    "segment",
    "status",
    "notes",
)


async def update_customer(
    db: AsyncSession,
    *,
    garage_id: int,
    actor_user_id: Optional[int],
    public_id: str,
    payload: CustomerUpdateRequest,
    request_id: Optional[str] = None,
    ip_address: Optional[str] = None,
) -> Customer:
    """Apply a partial update with optimistic locking and audit.

    Raises:
        CustomerNotFoundError: when the row does not exist in this garage.
        OptimisticLockError: when `payload.version` is stale.
        FiscalIdConflictError: when the new fiscal_id is taken by another live
            customer in the same garage.
    """
    logger.info("update_customer_start", public_id=public_id, garage_id=garage_id)
    target = await repo.get_by_public_id(
        db, garage_id, public_id, with_contacts=True
    )
    if target is None:
        raise CustomerNotFoundError(public_id)

    if target.version != payload.version:
        raise OptimisticLockError(expected=payload.version, actual=target.version)

    if payload.fiscal_id is not None:
        if target.type == CustomerType.company and not _SIRET_PATTERN.fullmatch(payload.fiscal_id):
            raise InvalidFiscalIdError(payload.fiscal_id)
        if payload.fiscal_id != target.fiscal_id:
            clash = await repo.find_by_fiscal_id(db, garage_id, payload.fiscal_id)
            if clash is not None and clash.id != target.id:
                raise FiscalIdConflictError(payload.fiscal_id)

    old_snapshot = _to_audit_dict(target)
    changed: dict[str, Any] = {}

    for field in _UPDATABLE_FIELDS:
        if field not in payload.model_fields_set:
            continue
        new_val = getattr(payload, field)
        if getattr(target, field) != new_val:
            setattr(target, field, new_val)
            changed[field] = new_val

    if "address" in payload.model_fields_set:
        new_addr = (
            payload.address.model_dump(exclude_none=True)
            if payload.address
            else None
        )
        if target.address != new_addr:
            target.address = new_addr
            changed["address"] = new_addr

    if not changed:
        # No-op update — still return the row but skip version bump and audit.
        return target

    target.version += 1
    target.updated_by_user_id = actor_user_id

    await audit(
        db=db,
        garage_id=garage_id,
        user_id=actor_user_id,
        component="customer",
        action="update",
        entity_type="customer",
        entity_id=target.id,
        old_value=old_snapshot,
        new_value=_to_audit_dict(target),
        request_id=request_id,
        ip_address=ip_address,
    )
    logger.info("update_customer_success", customer_id=target.id, garage_id=garage_id)
    return target


# ---------------------------------------------------------------------------
# Soft delete / restore
# ---------------------------------------------------------------------------


async def soft_delete_customer(
    db: AsyncSession,
    *,
    garage_id: int,
    actor_user_id: Optional[int],
    public_id: str,
    reason: Optional[str] = None,
    request_id: Optional[str] = None,
    ip_address: Optional[str] = None,
) -> Customer:
    logger.info("delete_customer_start", public_id=public_id, garage_id=garage_id)
    # include_deleted so the idempotent re-delete path can short-circuit.
    target = await repo.get_by_public_id(
        db, garage_id, public_id, include_deleted=True
    )
    if target is None:
        raise CustomerNotFoundError(public_id)
    if target.deleted_at is not None:
        return target  # already deleted, idempotent

    old_snapshot = _to_audit_dict(target)
    target.deleted_at = _utcnow()
    target.status = CustomerStatus.inactive
    target.updated_by_user_id = actor_user_id

    await audit(
        db=db,
        garage_id=garage_id,
        user_id=actor_user_id,
        component="customer",
        action="delete",
        entity_type="customer",
        entity_id=target.id,
        old_value=old_snapshot,
        new_value=_to_audit_dict(target),
        reason=reason,
        request_id=request_id,
        ip_address=ip_address,
    )
    logger.info("delete_customer_success", customer_id=target.id, garage_id=garage_id)
    return target


async def restore_customer(
    db: AsyncSession,
    *,
    garage_id: int,
    actor_user_id: Optional[int],
    public_id: str,
    request_id: Optional[str] = None,
    ip_address: Optional[str] = None,
) -> Customer:
    """Admin-only restore. Caller must enforce the role check.

    Re-checks fiscal_id uniqueness against live rows before restoring, since
    the partial unique index only constrains live rows.
    """
    logger.info("restore_customer_start", public_id=public_id, garage_id=garage_id)
    target = await repo.get_by_public_id(
        db, garage_id, public_id, include_deleted=True
    )
    if target is None:
        raise CustomerNotFoundError(public_id)
    if target.deleted_at is None:
        return target  # already live, idempotent

    if target.fiscal_id:
        clash = await repo.find_by_fiscal_id(db, garage_id, target.fiscal_id)
        if clash is not None and clash.id != target.id:
            raise FiscalIdConflictError(target.fiscal_id)

    old_snapshot = _to_audit_dict(target)
    target.deleted_at = None
    target.status = CustomerStatus.active
    target.updated_by_user_id = actor_user_id

    await audit(
        db=db,
        garage_id=garage_id,
        user_id=actor_user_id,
        component="customer",
        action="restore",
        entity_type="customer",
        entity_id=target.id,
        old_value=old_snapshot,
        new_value=_to_audit_dict(target),
        request_id=request_id,
        ip_address=ip_address,
    )
    logger.info("restore_customer_success", customer_id=target.id, garage_id=garage_id)
    return target


# ---------------------------------------------------------------------------
# Credit limit
# ---------------------------------------------------------------------------


async def update_credit_limit(
    db: AsyncSession,
    *,
    garage_id: int,
    actor_user_id: Optional[int],
    public_id: str,
    payload: CreditLimitUpdateRequest,
    request_id: Optional[str] = None,
    ip_address: Optional[str] = None,
) -> Customer:
    logger.info("update_credit_limit_start", public_id=public_id, garage_id=garage_id)
    target = await repo.get_by_public_id(db, garage_id, public_id)
    if target is None:
        raise CustomerNotFoundError(public_id)

    old_snapshot = _to_audit_dict(target)
    target.credit_limit_minor = payload.credit_limit_minor
    target.credit_currency = payload.credit_currency
    target.version += 1
    target.updated_by_user_id = actor_user_id

    await audit(
        db=db,
        garage_id=garage_id,
        user_id=actor_user_id,
        component="customer",
        action="credit_limit.update",
        entity_type="customer",
        entity_id=target.id,
        old_value=old_snapshot,
        new_value=_to_audit_dict(target),
        reason=payload.reason,
        request_id=request_id,
        ip_address=ip_address,
    )
    logger.info("update_credit_limit_success", customer_id=target.id, garage_id=garage_id)
    return target


# ---------------------------------------------------------------------------
# Balance ledger
# ---------------------------------------------------------------------------


async def add_balance_event(
    db: AsyncSession,
    *,
    garage_id: int,
    actor_user_id: Optional[int],
    public_id: str,
    payload: CustomerBalanceEventCreateRequest,
    request_id: Optional[str] = None,
    ip_address: Optional[str] = None,
) -> CustomerBalanceEvent:
    """Post a manual ledger event and update the customer's outstanding balance.

    Row-locks the customer to keep the ledger and outstanding column
    consistent under concurrent posts.
    """
    logger.info("add_balance_event_start", public_id=public_id, garage_id=garage_id)
    # SELECT ... FOR UPDATE to serialize concurrent balance adjustments.
    result = await db.execute(
        select(Customer)
        .where(
            Customer.garage_id == garage_id,
            Customer.deleted_at.is_(None),
        )
        .where(Customer.public_id == uuid_lib.UUID(public_id))
        .with_for_update()
    )
    target = result.scalar_one_or_none()
    if target is None:
        raise CustomerNotFoundError(public_id)

    event = CustomerBalanceEvent(
        customer_id=target.id,
        reason=payload.reason,
        delta_minor=payload.delta_minor,
        currency=payload.currency,
        created_by_user_id=actor_user_id,
    )
    db.add(event)

    target.outstanding_minor = target.outstanding_minor + payload.delta_minor
    target.updated_by_user_id = actor_user_id

    await db.flush()  # populate event.id

    await audit(
        db=db,
        garage_id=garage_id,
        user_id=actor_user_id,
        component="customer",
        action="balance.adjust",
        entity_type="customer_balance_event",
        entity_id=event.id,
        new_value={
            "customer_id": target.id,
            "reason": event.reason,
            "delta_minor": event.delta_minor,
            "currency": event.currency,
            "new_outstanding_minor": target.outstanding_minor,
        },
        reason=payload.note,
        request_id=request_id,
        ip_address=ip_address,
    )
    logger.info("add_balance_event_success", customer_id=target.id, event_id=event.id, garage_id=garage_id)
    return event


# ---------------------------------------------------------------------------
# Bulk import from Excel
# ---------------------------------------------------------------------------

# Tokens that appear in exported spreadsheets and must be treated as empty.
_EMPTY_TOKENS = {"", "undefined", "none", "null", "n/a", "na", "nan", "-", "—"}
_E164_RE = re.compile(r"^\+[1-9]\d{6,14}$")
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


@dataclass
class ImportedRow:
    """A single spreadsheet row normalised into customer fields."""

    type: CustomerType
    full_name: str
    company_name: Optional[str]
    phone_e164: Optional[str]
    email: Optional[str]
    address: Optional[dict[str, str]]


def _norm_header(value: Any) -> str:
    """Lower-case, accent-strip and trim a header cell for fuzzy matching."""
    import unicodedata

    s = unicodedata.normalize("NFKD", str(value or ""))
    s = s.encode("ascii", "ignore").decode("ascii")
    return s.lower().strip()


def _cell_text(value: Any) -> Optional[str]:
    """Return a trimmed string for a cell, or None for empty / placeholder values."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).strip()
    if s.lower() in _EMPTY_TOKENS:
        return None
    return s or None


def _clean_phone(value: Any) -> Optional[str]:
    """Normalise a phone cell to E.164, or None when it cannot be validated.

    Import is lenient: an unparseable number is dropped (customer still imported
    without a phone) rather than aborting the row.
    """
    s = _cell_text(value)
    if s is None:
        return None
    s = re.sub(r"[\s().\-/]", "", s)
    if s.startswith("00"):
        s = "+" + s[2:]
    if not s.startswith("+"):
        return None
    return s if _E164_RE.match(s) else None


def _clean_email(value: Any) -> Optional[str]:
    """Return a lower-cased valid email, or None when invalid / empty."""
    s = _cell_text(value)
    if s is None:
        return None
    s = s.lower()
    return s if _EMAIL_RE.match(s) else None


# Maps a normalised header to a logical field. Longer / more specific keys are
# checked first so that "nom de l'entreprise" wins over the bare "nom".
def _classify_header(header: str) -> Optional[str]:
    if not header:
        return None
    if "entreprise" in header or "societe" in header or "raison" in header or "nom de l" in header:
        return "company"
    if "prenom" in header:
        return "first_name"
    if header == "nom" or header.startswith("nom "):
        return "last_name"
    if "mail" in header:
        return "email"
    if "tel" in header or "phone" in header or "gsm" in header or "mobile" in header or "portable" in header:
        return "phone"
    if "code postal" in header or header in {"cp", "code"}:
        return "postal_code"
    if "ville" in header or "city" in header:
        return "city"
    if "adresse" in header or "address" in header or "rue" in header:
        return "street"
    return None


def parse_customers_xlsx(content: bytes) -> tuple[list[Optional[ImportedRow]], list[tuple[int, str]]]:
    """Parse an Excel workbook into normalised rows.

    Returns a tuple of:
      - a list where each entry is an `ImportedRow` (importable), or None for a
        fully-empty row that should be ignored silently;
      - a list of ``(row_number, reason)`` errors for rows that carried data but
        could not be turned into a valid customer.

    Raises ValueError when the file cannot be read as an Excel workbook.
    """
    import io

    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - surfaced as a 400 by the router
        raise ValueError(f"Fichier Excel illisible: {exc}") from exc

    ws = wb.active
    if ws is None:
        raise ValueError("Le classeur ne contient aucune feuille.")

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], []

    # Build column index -> logical field mapping from the header row.
    field_by_col: dict[int, str] = {}
    for idx, raw in enumerate(header_row):
        field = _classify_header(_norm_header(raw))
        if field is not None and idx not in field_by_col.values():
            field_by_col[idx] = field

    parsed: list[Optional[ImportedRow]] = []
    errors: list[tuple[int, str]] = []

    for offset, raw_row in enumerate(rows_iter):
        row_number = offset + 2  # header is row 1
        values: dict[str, Any] = {}
        for idx, field in field_by_col.items():
            values[field] = raw_row[idx] if idx < len(raw_row) else None

        last_name = _cell_text(values.get("last_name"))
        first_name = _cell_text(values.get("first_name"))
        company = _cell_text(values.get("company"))
        phone = _clean_phone(values.get("phone"))
        email = _clean_email(values.get("email"))
        street = _cell_text(values.get("street"))
        city = _cell_text(values.get("city"))
        postal_code = _cell_text(values.get("postal_code"))

        person = " ".join(p for p in (first_name, last_name) if p).strip()

        # Fully empty row: ignore without counting it.
        if not any([person, company, phone, email, street, city, postal_code]):
            parsed.append(None)
            continue

        if not person and not company:
            errors.append((row_number, "Nom et prénom manquants."))
            parsed.append(None)
            continue

        if company:
            ctype = CustomerType.company
            full_name = (person or company)[:255]
            company_name = company[:255]
        else:
            ctype = CustomerType.individual
            full_name = person[:255]
            company_name = None

        address: dict[str, str] = {}
        if street:
            address["street"] = street[:255]
        if city:
            address["city"] = city[:100]
        if postal_code:
            address["postal_code"] = postal_code[:20]

        parsed.append(
            ImportedRow(
                type=ctype,
                full_name=full_name,
                company_name=company_name,
                phone_e164=phone,
                email=email,
                address=address or None,
            )
        )

    wb.close()
    return parsed, errors


async def import_customers(
    db: AsyncSession,
    *,
    garage_id: int,
    actor_user_id: Optional[int],
    content: bytes,
    request_id: Optional[str] = None,
    ip_address: Optional[str] = None,
):
    """Seed customers from an uploaded Excel file.

    Parsing is lenient — empty / malformed optional cells are dropped rather
    than aborting the whole import. Rows that duplicate an existing customer
    are skipped. Duplicates are detected against both existing customers and
    earlier rows in the same file, using the strongest available key:
    phone → email → (full_name [+ company_name]) when neither contact exists.
    All valid rows are inserted in a single transaction; the caller owns the
    commit.
    """
    from app.modules.customer.schemas import CustomerImportResult, CustomerImportRowError

    parsed, parse_errors = parse_customers_xlsx(content)

    errors: list[CustomerImportRowError] = [
        CustomerImportRowError(row=r, reason=msg) for r, msg in parse_errors
    ]
    created = 0
    skipped = 0
    total_rows = 0

    seen_phones: set[str] = set()
    seen_emails: set[str] = set()
    seen_names: set[tuple[str, str]] = set()

    for offset, row in enumerate(parsed):
        if row is None:
            continue
        total_rows += 1
        row_number = offset + 2

        # De-dupe within the file and against existing customers, using the
        # strongest identity key the row carries.
        if row.phone_e164 and (
            row.phone_e164 in seen_phones
            or await repo.find_by_phone(db, garage_id, row.phone_e164) is not None
        ):
            skipped += 1
            continue
        if row.email and (
            row.email in seen_emails
            or await repo.find_by_email(db, garage_id, row.email) is not None
        ):
            skipped += 1
            continue
        # Fallback for rows with no phone and no email: match on name so a
        # re-import of the same list does not create duplicates.
        if not row.phone_e164 and not row.email:
            name_key = (row.full_name.lower(), (row.company_name or "").lower())
            if name_key in seen_names or (
                await repo.find_by_name(db, garage_id, row.full_name, row.company_name)
                is not None
            ):
                skipped += 1
                continue
            seen_names.add(name_key)

        customer = Customer(
            garage_id=garage_id,
            type=row.type,
            full_name=row.full_name,
            company_name=row.company_name,
            phone_e164=row.phone_e164,
            email=row.email,
            address=row.address,
            created_by_user_id=actor_user_id,
            updated_by_user_id=actor_user_id,
        )
        db.add(customer)
        created += 1
        if row.phone_e164:
            seen_phones.add(row.phone_e164)
        if row.email:
            seen_emails.add(row.email)

    await db.flush()

    await audit(
        db=db,
        garage_id=garage_id,
        user_id=actor_user_id,
        component="customer",
        action="import",
        entity_type="customer",
        entity_id=None,
        new_value={
            "total_rows": total_rows,
            "created": created,
            "skipped": skipped,
            "errors": len(errors),
        },
        request_id=request_id,
        ip_address=ip_address,
    )
    logger.info(
        "import_customers_done",
        garage_id=garage_id,
        total_rows=total_rows,
        created=created,
        skipped=skipped,
        errors=len(errors),
    )
    return CustomerImportResult(
        total_rows=total_rows,
        created=created,
        skipped=skipped,
        errors=errors,
    )
